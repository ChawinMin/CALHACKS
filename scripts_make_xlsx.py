"""Build wound_dataset.xlsx: one row per image, labeled by its class folder, with an embedded thumbnail."""
import io
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image

ROOT = Path(__file__).resolve().parent / "dataset" / "Wound_dataset"
OUT = Path(__file__).resolve().parent / "dataset" / "wound_dataset.xlsx"
THUMB = 120  # px, max side


def natural(p: Path):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", p.name)]


wb = Workbook()
ws = wb.active
ws.title = "Images"
headers = ["#", "Image", "Wound Type", "File Name", "Relative Path", "Width (px)", "Height (px)"]
ws.append(headers)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="305496")
    c.alignment = Alignment(horizontal="center", vertical="center")
ws.freeze_panes = "A2"

counts = {}
row = 2
n = 0
for cls in sorted(p for p in ROOT.iterdir() if p.is_dir()):
    label = cls.name.replace("_", " ")
    for img in sorted((f for f in cls.iterdir() if f.suffix.lower() in {".jpg", ".jpeg", ".png"}), key=natural):
        n += 1
        with Image.open(img) as im:
            im = im.convert("RGB")
            w, h = im.size
            im.thumbnail((THUMB, THUMB))
            buf = io.BytesIO()
            im.save(buf, "JPEG", quality=80)
            tw, th = im.size
        buf.seek(0)
        xi = XLImage(buf)
        xi.width, xi.height = tw, th
        ws.add_image(xi, f"B{row}")
        ws.append([n, None, label, img.name, str(img.relative_to(ROOT.parent)), w, h])
        ws.row_dimensions[row].height = THUMB * 0.76 + 4
        for col in "ACDEFG":
            ws[f"{col}{row}"].alignment = Alignment(vertical="center")
        counts[label] = counts.get(label, 0) + 1
        row += 1

for col, width in zip("ABCDEFG", [6, 19, 16, 24, 48, 11, 11]):
    ws.column_dimensions[col].width = width
ws.auto_filter.ref = f"A1:G{row - 1}"

s = wb.create_sheet("Summary")
s.append(["Wound Type", "Image Count"])
for c in s[1]:
    c.font = Font(bold=True)
for k, v in counts.items():
    s.append([k, v])
s.append(["Total", f"=SUM(B2:B{len(counts) + 1})"])
s[f"A{len(counts) + 2}"].font = Font(bold=True)
s.column_dimensions["A"].width = 18

wb.save(OUT)
print(n, "images ->", OUT, counts)
